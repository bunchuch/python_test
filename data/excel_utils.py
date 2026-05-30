from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def apply_excel_styles(output: BytesIO, df: pd.DataFrame) -> BytesIO:
    """
    Apply formatting to an already-written Excel workbook.

    Strategy for large datasets (7+ days):
      - Style header row cell-by-cell  (only 44 cells — always fast)
      - Register an openpyxl Table with a built-in stripe style for data rows
        This avoids an O(rows × cols) cell loop that would take 30-60 s on
        thousands of rows.
    """
    output.seek(0)
    wb = load_workbook(output)
    ws = wb.active

    n_cols    = len(df.columns)
    n_rows    = ws.max_row          # includes header
    last_col  = get_column_letter(n_cols)

    # ── Header row (44 cells — fast) ──────────────────────────────────────────
    HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
    CENTER      = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, n_cols + 1):
        cell            = ws.cell(row=1, column=col_idx)
        cell.fill       = HEADER_FILL
        cell.font       = HEADER_FONT
        cell.alignment  = CENTER

    # ── Data rows — use a Table style (no per-cell loop) ─────────────────────
    # TableStyleMedium9 = blue header + light-blue alternating stripes,
    # which matches the #1F4E79 / #D6E4F0 palette.
    table_ref = f"A1:{last_col}{n_rows}"
    tbl = Table(displayName="SabreData", ref=table_ref)
    tbl.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showRowStripes=True,
        showColumnStripes=False,
        showFirstColumn=False,
        showLastColumn=False,
    )
    ws.add_table(tbl)

    # ── Auto-fit column widths (capped at 40) ─────────────────────────────────
    for col_idx, col_name in enumerate(df.columns, start=1):
        col_letter = get_column_letter(col_idx)
        header_len = len(str(col_name))
        max_data   = (
            df.iloc[:, col_idx - 1].astype(str).str.len().max()
            if len(df) > 0 else 0
        )
        ws.column_dimensions[col_letter].width = min(
            max(header_len, max_data, 8) + 2, 40
        )

    ws.freeze_panes = "A2"

    styled = BytesIO()
    wb.save(styled)
    styled.seek(0)
    return styled
